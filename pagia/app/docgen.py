"""Παραγωγή εγγράφων: πρακτικά, διαβιβαστικό, ημερολόγιο κίνησης.

Η αντικατάσταση μεταβλητών στα .docx γίνεται με προσοχή στα «runs»:
το Word σπάει συχνά μια λέξη σε πολλά κομμάτια (π.χ. λόγω ορθογραφικού
ελέγχου), οπότε η αναζήτηση κειμένου ανά run αποτυγχάνει. Εδώ ενώνουμε
τα runs κάθε παραγράφου, κάνουμε την αντικατάσταση, και ξαναγράφουμε
το αποτέλεσμα στο πρώτο run διατηρώντας τη μορφοποίησή του.
"""

import io
import os
import re
import unicodedata
from datetime import date

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

from db import (DAYS_GR, DEFAULT_CATEGORY, MONTHS_ACC, MONTHS_GEN,
                MONTHS_NOM, PRAKTIKO_TYPES)

TEMPLATE_DIR = os.environ.get("TEMPLATE_DIR", "/data/templates")

# Τα πρότυπα που διαχειρίζεται η εφαρμογή. Κλειδί -> (αρχείο, περιγραφή)
TEMPLATES = {
    "ylikwn": ("ΠΡΟΤΥΠΟ_-_Πρακτικό_Παραλαβής_Υλικών.docx",
               "Πρακτικό Παραλαβής Υλικών"),
    "ekpaideysis": ("ΠΡΟΤΥΠΟ_-_Πρακτικό_Παραλαβής_Υλικών_Εκπαίδευσης.docx",
                    "Πρακτικό Παραλαβής Υλικών Εκπαίδευσης"),
    "kalis": ("ΠΡΟΤΥΠΟ_-_Πρακτικό_Καλής_Εκτέλεσης.docx",
              "Πρακτικό Καλής Εκτέλεσης"),
    "diavivastiko": ("ΠΡΟΤΥΠΟ_-_Διαβιβαστικό.docx", "Διαβιβαστικό"),
}

# Μεταβλητές που γεμίζει η εφαρμογή. Ό,τι άλλο μπει σε πρότυπο μένει ως έχει.
COMMON_FIELDS = ["ΣΧΟΛΗ", "ΣΧΟΛΗ_ΓΕΝΙΚΗ", "ΤΟΠΟΣ", "ΑΠΟΦΑΣΗ",
                 "ΕΠΙΤΡΟΠΗ_Α", "ΕΠΙΤΡΟΠΗ_Β", "ΕΠΙΤΡΟΠΗ_Γ",
                 "ΑΡ_ΠΡΩΤΟΚΟΛΛΟΥ", "ΗΜΕΡΟΜΗΝΙΑ_ΠΡΩΤΟΚΟΛΛΟΥ"]
PRAKTIKO_FIELDS = COMMON_FIELDS + [
    "ΗΜΕΡΑ", "ΩΡΑ", "ΕΠΩΝΥΜΙΑ_ΠΡΟΜΗΘΕΥΤΗ", "ΑΦΜ", "ΑΡΙΘΜΟΣ_ΠΑΡΑΣΤΑΤΙΚΟΥ",
    "ΗΜΕΡΟΜΗΝΙΑ_ΠΑΡΑΣΤΑΤΙΚΟΥ", "ΜΕΙΚΤΗ_ΑΞΙΑ", "ΕΙΔΟΣ_ΠΑΡΑΣΤΑΤΙΚΟΥ",
    "ΕΙΔΙΚΟΤΗΤΑ", "ΕΞΑΜΗΝΟ", "ΤΜΗΜΑ"]
DIAVIVASTIKO_FIELDS = COMMON_FIELDS + [
    "ΜΗΝΑΣ", "ΜΗΝΑΣ_ΓΕΝΙΚΗ", "ΜΗΝΑΣ_ΑΙΤΙΑΤΙΚΗ", "ΜΗΝΑΣ_ΟΝΟΜΑΣΤΙΚΗ",
    "ΕΤΟΣ", "ΛΙΣΤΑ_ΠΑΡΑΣΤΑΤΙΚΩΝ"]


def known_fields(key: str) -> list:
    return DIAVIVASTIKO_FIELDS if key == "diavivastiko" else PRAKTIKO_FIELDS
PLACEHOLDER_RE = re.compile(r"«([A-Za-zΑ-Ωα-ωΆ-Ώά-ώ_0-9]+)»")


def gr_upper(text: str) -> str:
    """Κεφαλαία χωρίς τόνους, όπως γράφονται στα έγγραφα.

    Το σκέτο .upper() δίνει «ΙΩΆΝΝΗΣ». Στα ελληνικά κεφαλαία ο τόνος
    φεύγει, τα διαλυτικά όμως μένουν: Αϊβαλής → ΑΪΒΑΛΗΣ.
    """
    TONOS = "\u0301"        # οξεία
    up = unicodedata.normalize("NFD", (text or "").upper())
    return unicodedata.normalize("NFC", up.replace(TONOS, ""))


def _fmt_date(d):
    return d.strftime("%d/%m/%Y") if d else ""


def _fmt_money(v):
    """Ελληνική μορφή: τελεία για χιλιάδες, κόμμα για δεκαδικά."""
    s = f"{float(v or 0):,.2f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def _set_multiline(par, text):
    """Γράφει πολλές γραμμές σε μία παράγραφο με πραγματικά <w:br/>.

    Το σκέτο "\n" μέσα σε run ΔΕΝ εμφανίζεται ως αλλαγή γραμμής στο Word."""
    if not par.runs:
        return
    first = par.runs[0]
    for r in par.runs[1:]:
        r._element.getparent().remove(r._element)
    first.text = ""
    for i, line in enumerate(text.split("\n")):
        if i:
            first.add_break()
        if line:
            first.add_text(line)


def _replace_in_paragraph(par, mapping):
    full = "".join(r.text for r in par.runs)
    if "«" not in full:
        return
    # Πολυγραμμική τιμή που καταλαμβάνει μόνη της την παράγραφο
    for key, val in mapping.items():
        if "\n" in str(val) and full.strip() == f"«{key}»":
            _set_multiline(par, str(val))
            return
    new = full
    for key, val in mapping.items():
        new = new.replace(f"«{key}»", str(val))
    if new == full:
        return
    # Το πρώτο run κρατά όλο το κείμενο· τα υπόλοιπα αδειάζουν.
    if par.runs:
        par.runs[0].text = new
        for r in par.runs[1:]:
            r.text = ""


def _walk(doc):
    """Όλες οι παράγραφοι, μαζί με πίνακες, κεφαλίδες και υποσέλιδα."""
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for it in cell.tables:
                    for r2 in it.rows:
                        for c2 in r2.cells:
                            yield from c2.paragraphs
    for sec in doc.sections:
        for part in (sec.header, sec.footer,
                     sec.first_page_header, sec.first_page_footer,
                     sec.even_page_header, sec.even_page_footer):
            if part is None:
                continue
            yield from part.paragraphs
            for t in part.tables:
                for row in t.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def fill_template(template_name: str, mapping: dict,
                  one_page: bool = False) -> bytes:
    path = os.path.join(TEMPLATE_DIR, template_name)
    if not os.path.exists(path):
        raise FileNotFoundError(f"Λείπει το πρότυπο: {template_name}")
    doc = Document(path)
    for par in _walk(doc):
        _replace_in_paragraph(par, mapping)
    if one_page:
        compact(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def placeholders_in(source) -> set:
    """Ποιες μεταβλητές «ΧΧΧ» περιέχει ένα πρότυπο.

    Δέχεται διαδρομή αρχείου ή bytes, ώστε να ελέγχεται και αρχείο που
    μόλις ανέβηκε, πριν αποθηκευτεί."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    elif not os.path.exists(source):
        return set()
    doc = Document(source)
    found = set()
    for par in _walk(doc):
        full = "".join(r.text for r in par.runs)
        found.update(PLACEHOLDER_RE.findall(full))
    return found


def _is_truly_empty(par) -> bool:
    """Κενή παράγραφος = χωρίς κείμενο ΚΑΙ χωρίς περιεχόμενο.

    Το `par.text` επιστρέφει κενό και όταν η παράγραφος περιέχει μόνο
    εικόνες. Στα πρότυπα η τελευταία παράγραφος κρατά τα λογότυπα και
    τη ζώνη υπογραφών — ο έλεγχος πρέπει να τα δει.
    """
    if par.text.strip():
        return False
    el = par._element
    for tag in ("w:drawing", "w:pict", "w:object", "w:txbxContent",
                "w:fldSimple", "w:instrText", "w:tab"):
        if el.findall(".//" + qn(tag)):
            return False
    return True


def compact(doc, space=2.0, line=1.10, bottom_cm=0.8, footer_cm=0.6):
    """Κερδίζει χώρο ώστε το πρακτικό να χωρά άνετα σε μία σελίδα.

    Μετρημένο στο πρότυπο: το κείμενο τελείωνε 2,1 εκ. πάνω από το κάτω
    περιθώριο — πολύ λίγο, γι' αυτό το Word το έσπρωχνε σε δεύτερη
    σελίδα. Με τις ρυθμίσεις εδώ τα ελεύθερα γίνονται 3,7 εκ.

    Τι δουλεύει και τι όχι:
      * κάτω περιθώριο 1,8 → 0,8 εκ.  →  +1,0 εκ.
      * μάζεμα κενών και διάστιχου     →  +0,6 εκ.
      * πλάγια περιθώρια               →  καμία διαφορά, το κείμενο
                                          απλώς απλώνει σε λιγότερες
                                          γραμμές· δεν τα πειράζουμε

    Τα μεγέθη γραμματοσειράς ΔΕΝ αλλάζουν, και οι παράγραφοι που
    κρατούν εικόνες ή στηλοθέτες ΔΕΝ αφαιρούνται.
    """
    for sec in doc.sections:
        sec.bottom_margin = Cm(bottom_cm)
        sec.footer_distance = Cm(footer_cm)

    while doc.paragraphs and _is_truly_empty(doc.paragraphs[-1]):
        el = doc.paragraphs[-1]._element
        el.getparent().remove(el)

    for par in doc.paragraphs:
        pf = par.paragraph_format
        pf.space_before = Pt(space)
        pf.space_after = Pt(space)
        pf.line_spacing = line
    return doc


# ----------------------------------------------------------------- ΠΡΑΚΤΙΚΑ

def praktiko_mapping(inv, org: dict = None) -> dict:
    d = inv.protocol_date or date.today()
    dep = inv.department
    return {
        **(org or {}),
        "ΕΙΔΟΣ_ΠΑΡΑΣΤΑΤΙΚΟΥ": inv.doc_type.name if inv.doc_type else "",
        "ΑΡ_ΠΡΩΤΟΚΟΛΛΟΥ": inv.protocol_ref,
        "ΗΜΕΡΟΜΗΝΙΑ_ΠΡΩΤΟΚΟΛΛΟΥ": _fmt_date(d),
        "ΗΜΕΡΑ": DAYS_GR[d.weekday()],
        "ΕΠΩΝΥΜΙΑ_ΠΡΟΜΗΘΕΥΤΗ": inv.supplier.name,
        "ΑΦΜ": inv.supplier.afm,
        "ΑΡΙΘΜΟΣ_ΠΑΡΑΣΤΑΤΙΚΟΥ": inv.number,
        "ΗΜΕΡΟΜΗΝΙΑ_ΠΑΡΑΣΤΑΤΙΚΟΥ": _fmt_date(inv.doc_date),
        "ΜΕΙΚΤΗ_ΑΞΙΑ": _fmt_money(inv.gross),
        "ΕΙΔΙΚΟΤΗΤΑ": dep.specialty if dep else "",
        "ΕΞΑΜΗΝΟ": dep.semester if dep else "",
        "ΤΜΗΜΑ": dep.section if dep else "",
    }


def build_praktiko(inv, org: dict = None) -> tuple[str, bytes]:
    ptype = inv.praktiko_type
    if not ptype:
        raise ValueError("Η δαπάνη αυτή δεν παράγει πρακτικό.")
    label, template = PRAKTIKO_TYPES[ptype][0], PRAKTIKO_TYPES[ptype][1]
    data = fill_template(template, praktiko_mapping(inv, org), one_page=True)
    safe = re.sub(r"[^\w\-. ]", "", inv.supplier.name)[:40].strip()
    num = inv.protocol_number or "ΧΩΡΙΣ-ΠΡΩΤ"
    fname = f"{num}_Πρακτικό {label}_{safe}.docx"
    return fname, data


# ------------------------------------------------------------- ΔΙΑΒΙΒΑΣΤΙΚΟ

def diavivastiko_block(n: int, inv) -> str:
    """Ένα αριθμημένο μπλοκ του διαβιβαστικού.

    Δομή:
      N. Έντυπα Δαπάνης Προμήθειας <κατηγορία> μέσω Παγίας Δαπάνης:
         προμηθευτή <επωνυμία>
      Α. <αριθμός παραστατικού>/<ημερομηνία>
      Β. <πρακτικό>                      — αν η δαπάνη παράγει πρακτικό
      Γ. Αποδεικτικό εξόφλησης ...       — αν ο τρόπος πληρωμής το απαιτεί
    """
    category = inv.praktiko_category

    # Ο αριθμός καταχωρείται με το πρόθεμα του είδους («ΤΠΔΑ 25946»),
    # οπότε δεν ξαναγράφεται η ονομασία του είδους μπροστά.
    lines = [f"{n}. Έντυπα Δαπάνης Προμήθειας {category} μέσω Παγίας "
             f"Δαπάνης: προμηθευτή {inv.supplier.name}",
             f"Α. {inv.number}/{inv.doc_date.strftime('%d-%m-%Y')}"]

    praktiko = inv.praktiko_line
    if praktiko:
        lines.append(f"Β. {praktiko}")

    pm = inv.payment_method
    # Χωρίς δηλωμένο τρόπο πληρωμής θεωρούμε ότι χρειάζεται αποδεικτικό:
    # καλύτερα να το σβήσεις εσύ παρά να λείπει.
    if pm is None or pm.needs_proof:
        lines.append("Γ. Αποδεικτικό εξόφλησης μέσω τραπεζικής μεταφοράς")
    return "\n".join(lines)


def build_diavivastiko(batch, invoices, org: dict = None) -> tuple[str, bytes]:
    lista = "\n".join(diavivastiko_block(i, inv)
                      for i, inv in enumerate(invoices, 1))
    d = batch.protocol_date or date.today()
    m, y = batch.month, batch.year
    mapping = {
        **(org or {}),
        "ΑΡ_ΠΡΩΤΟΚΟΛΛΟΥ": batch.protocol_ref,
        "ΗΜΕΡΟΜΗΝΙΑ_ΠΡΩΤΟΚΟΛΛΟΥ": _fmt_date(d),
        # Στο πρότυπο το «ΜΗΝΑΣ» χρησιμοποιείται σε γενική, με το έτος:
        # «Αποστολή Τιμολογίων Ιουνίου 2026»
        "ΜΗΝΑΣ": f"{MONTHS_GEN[m - 1]} {y}",
        "ΜΗΝΑΣ_ΓΕΝΙΚΗ": f"{MONTHS_GEN[m - 1]} {y}",
        "ΜΗΝΑΣ_ΑΙΤΙΑΤΙΚΗ": f"{MONTHS_ACC[m - 1]} {y}",
        "ΜΗΝΑΣ_ΟΝΟΜΑΣΤΙΚΗ": MONTHS_NOM[m - 1],
        "ΕΤΟΣ": y,
        "ΛΙΣΤΑ_ΠΑΡΑΣΤΑΤΙΚΩΝ": lista,
    }
    data = fill_template("ΠΡΟΤΥΠΟ_-_Διαβιβαστικό.docx", mapping)
    fname = f"Διαβιβαστικό_{MONTHS_GEN[m - 1]}_{y}.docx"
    return fname, data


# ---------------------------------------------------------------- ΗΜΕΡΟΛΟΓΙΟ

HDR = ["Α/Α", "ΧΡΗΣΗ ΕΤΟΥΣ", "ΗΜΕΡΟΜΗΝΙΑ ΚΙΝΗΣΗΣ",
       "ΑΙΤΙΟΛΟΓΙΑ ΚΙΝΗΣΗΣ / ΕΠΩΝΥΜΙΑ", "ΑΦΜ", "ΑΡΙΘΜΟΣ ΠΑΡΑΣΤΑΤΙΚΟΥ",
       "ΕΙΔΟΣ ΠΑΡΑΣΤΑΤΙΚΟΥ", "ΚΩΔ. ΔΑΠΑΝΗΣ / ΟΝΟΜΑΣΙΑ", "ΤΡΟΠΟΣ ΚΙΝΗΣΗΣ",
       "ΚΑΤΗΓΟΡΙΑ ΕΞΟΔΩΝ", "ΠΙΣΤΩΣΗ", "ΧΡΕΩΣΗ", "ΗΜΕΡΟΜΗΝΙΑ ΠΑΡΑΣΤΑΤΙΚΟΥ"]


def build_imerologio(year, month, invoices, org: dict = None) -> tuple[str, bytes]:
    """Ημερολόγιο κίνησης, έτοιμο για εκτύπωση σε οριζόντιο Α4."""
    org = org or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Ημερολόγιο"

    title = (f"ΗΜΕΡΟΛΟΓΙΟ ΚΙΝΗΣΗΣ ΠΑΓΙΑΣ ΠΡΟΚΑΤΑΒΟΛΗΣ — "
             f"{MONTHS_NOM[month - 1]} {year}")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR))
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=13)
    c.alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="DDE5EE")

    for j, h in enumerate(HDR, 1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = Font(bold=True, size=9)
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="center",
                                   horizontal="center")

    row = 4
    total = 0.0
    for n, inv in enumerate(invoices, 1):
        vals = [n, year, inv.payment_date or inv.doc_date, inv.supplier.name,
                inv.supplier.afm, inv.number,
                inv.doc_type.name if inv.doc_type else "",
                inv.expense_code.label if inv.expense_code else "",
                inv.payment_method.name if inv.payment_method else "",
                "Αγορά", None, round(inv.gross, 2), inv.doc_date]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.border = border
            cell.font = Font(size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if j in (3, 13) and v:
                cell.number_format = "DD/MM/YYYY"
            if j in (11, 12):
                cell.number_format = '#,##0.00 "€"'
        total += inv.gross
        row += 1

    ws.cell(row=row, column=11, value="ΣΥΝΟΛΟ").font = Font(bold=True)
    tc = ws.cell(row=row, column=12, value=round(total, 2))
    tc.font = Font(bold=True)
    tc.number_format = '#,##0.00 "€"'
    tc.border = border
    last_data_row = row

    # ---------------------------------------------------- ζώνη υπογραφής
    def centered(r, text, **font):
        ws.merge_cells(start_row=r, start_column=1,
                       end_row=r, end_column=len(HDR))
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = Font(**font)
        cell.alignment = Alignment(horizontal="center")
        return r + 1

    r = last_data_row + 3
    r = centered(r, "ΘΕΩΡΗΘΗΚΕ", bold=True, size=11)
    r = centered(r, f"{org.get('ΤΟΠΟΣ', '')}, "
                    f"{MONTHS_NOM[month - 1]} {year}", bold=True, size=10)
    r += 1
    r = centered(r, gr_upper(org.get("ΤΙΤΛΟΣ_ΔΙΕΥΘΥΝΤΗ", "")), bold=True, size=10)
    school = org.get("ΣΧΟΛΗ_ΓΕΝΙΚΗ") or org.get("ΣΧΟΛΗ", "")
    r = centered(r, school, size=10)
    r += 2
    r = centered(r, gr_upper(org.get("ΔΙΕΥΘΥΝΤΗΣ", "")), bold=True, size=11)

    widths = [5, 9, 13, 30, 11, 17, 18, 26, 20, 12, 11, 11, 13]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"

    # ------------------------------------------------ ρυθμίσεις εκτύπωσης
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9          # A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1     # όλες οι στήλες σε ένα πλάτος σελίδας
    ws.page_setup.fitToHeight = 0    # όσες σελίδες χρειαστούν καθ' ύψος
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5,
                                  header=0.2, footer=0.2)
    ws.print_title_rows = "3:3"      # η κεφαλίδα επαναλαμβάνεται σε κάθε σελίδα
    ws.print_area = f"A1:{get_column_letter(len(HDR))}{r}"
    ws.oddFooter.center.text = "Σελίδα &P από &N"
    ws.oddFooter.center.size = 8
    ws.print_options.horizontalCentered = True

    buf = io.BytesIO()
    wb.save(buf)
    fname = f"Ημερολόγιο_Κίνησης_{MONTHS_NOM[month - 1]}_{year}.xlsx"
    return fname, buf.getvalue()


# --------------------------------------------------------- ΣΥΓΚΕΝΤΡΩΤΙΚΑ

def _sheet(ws, title, headers, rows, widths):
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="DDE5EE")

    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(headers))
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=12)
    c.alignment = Alignment(horizontal="center")

    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = Font(bold=True, size=9)
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, horizontal="center")

    r = 4
    totals = [0] * len(headers)
    # Ποιες στήλες αθροίζονται. Χωρίς αυτό, μηδενικό σύνολο (π.χ. ΦΠΑ)
    # αφήνει το κελί κενό αντί για 0,00 €.
    numeric = [False] * len(headers)
    for row in rows:
        for j, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = border
            cell.font = Font(size=9)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                numeric[j - 1] = True
                totals[j - 1] += v
                if isinstance(v, float):
                    cell.number_format = '#,##0.00 "€"'
        r += 1

    tc = ws.cell(row=r, column=1, value="ΣΥΝΟΛΟ")
    tc.font = Font(bold=True)
    tc.border = border
    for j, t in enumerate(totals, 1):
        if not numeric[j - 1]:
            continue
        cell = ws.cell(row=r, column=j, value=round(t, 2))
        cell.font = Font(bold=True)
        cell.border = border
        if isinstance(t, float):
            cell.number_format = '#,##0.00 "€"'
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"


def build_summary(period: str, by_code: dict, by_dept: dict) -> tuple[str, bytes]:
    """Συγκεντρωτικά ανά κωδικό δαπάνης και ανά τμήμα."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ανά κωδικό δαπάνης"
    _sheet(ws, f"ΣΥΝΟΛΑ ΑΝΑ ΚΩΔΙΚΟ ΔΑΠΑΝΗΣ — {period}",
           ["Κωδικός δαπάνης", "Κωδικός", "Πλήθος",
            "Καθαρή αξία", "ΦΠΑ", "Μεικτή αξία"],
           [[k, v[0], v[1], round(v[2], 2), round(v[3], 2), round(v[4], 2)]
            for k, v in sorted(by_code.items(), key=lambda kv: -kv[1][4])],
           [52, 20, 10, 14, 14, 14])

    ws2 = wb.create_sheet("Ανά τμήμα")
    _sheet(ws2, f"ΥΛΙΚΑ ΕΚΠΑΙΔΕΥΣΗΣ ΑΝΑ ΤΜΗΜΑ — {period}",
           ["Τμήμα", "Πλήθος", "Καθαρή αξία", "ΦΠΑ", "Μεικτή αξία"],
           [[k, v[0], round(v[1], 2), round(v[2], 2), round(v[3], 2)]
            for k, v in sorted(by_dept.items(), key=lambda kv: -kv[1][3])],
           [52, 10, 14, 14, 14])

    buf = io.BytesIO()
    wb.save(buf)
    safe = period.replace("/", "-").replace(" ", "_")
    return f"Συγκεντρωτικά_{safe}.xlsx", buf.getvalue()
